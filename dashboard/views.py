import json
from datetime import datetime
import pandas as pd
from django.contrib import messages
from django.db import transaction
from django.shortcuts import render, redirect
from django.utils.dateparse import parse_date
from .forms import ExcelUploadForm
from .models import Stock, WeldingNeed, DocumentationRecord
import csv
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.shortcuts import render, redirect, get_object_or_404
from urllib.parse import urlencode
from .forms import (
    ExcelUploadForm,
    StockEditForm,
    WeldingNeedEditForm,
    DocumentationEditForm,
)
from collections import Counter


def upload_excel(request):
    # Существующая загрузка Excel остаётся доступной.
    if request.method == "POST":
        return _process_excel_upload(request)

    return render(
        request,
        "dashboard/upload.html",
        {"form": ExcelUploadForm()},
    )


def _process_excel_upload(request):
    if request.method == "POST":
        form = ExcelUploadForm(request.POST, request.FILES)

        if form.is_valid():
            data_type = form.cleaned_data["data_type"]
            file = form.cleaned_data["file"]

            try:
                df = pd.read_excel(file, sheet_name=0)
                df.columns = df.columns.astype(str).str.strip()

                if df.empty:
                    raise ValueError("Файл не содержит записей.")

                if df.columns.duplicated().any():
                    raise ValueError(
                        "В файле есть повторяющиеся названия столбцов."
                    )

                required_columns = {
                    "stock": {
                        "record_id", "date", "batch_id",
                        "employee", "plan_qty", "fact_qty",
                    },
                    "welding": {
                        "record_id", "date", "batch_id", "need_qty",
                    },
                    "documentation": {
                        "record_id", "date", "employee",
                        "plan_docs", "fact_docs",
                    },
                }[data_type]

                missing = required_columns - set(df.columns)

                if missing:
                    raise ValueError(
                        "Не найдены столбцы: "
                        + ", ".join(sorted(missing))
                    )

                df = df[list(required_columns)].copy()

                # Проверяем обязательные поля и преобразуем даты.
                date_values = pd.to_datetime(
                    df["date"], errors="coerce"
                )

                if date_values.isna().any():
                    raise ValueError(
                        "В столбце date есть пустые или неверные даты."
                    )

                df["date"] = date_values.dt.date

                if df["record_id"].isna().any():
                    raise ValueError(
                        "У некоторых записей отсутствует record_id."
                    )

                if df["record_id"].astype(str).str.strip().eq("").any():
                    raise ValueError("Обнаружен пустой record_id.")

                with transaction.atomic():
                    if data_type == "stock":
                        records = []

                        for row in df.to_dict("records"):
                            plan = pd.to_numeric(
                                row["plan_qty"], errors="coerce"
                            )
                            fact = pd.to_numeric(
                                row["fact_qty"], errors="coerce"
                            )

                            if pd.isna(plan) or plan % 1 != 0:
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "plan_qty должен быть целым числом."
                                )

                            if not pd.isna(fact) and fact % 1 != 0:
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "fact_qty должен быть целым числом."
                                )

                            if pd.isna(row["batch_id"]) or pd.isna(
                                row["employee"]
                            ):
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "batch_id и employee обязательны."
                                )

                            records.append(
                                Stock(
                                    record_id=str(row["record_id"]).strip(),
                                    date=row["date"],
                                    batch_id=str(row["batch_id"]).strip(),
                                    employee=str(row["employee"]).strip(),
                                    plan_qty=int(plan),
                                    fact_qty=(
                                        None if pd.isna(fact) else int(fact)
                                    ),
                                )
                            )

                        Stock.objects.bulk_create(records)

                    elif data_type == "welding":
                        records = []

                        for row in df.to_dict("records"):
                            qty = pd.to_numeric(
                                row["need_qty"], errors="coerce"
                            )

                            if pd.isna(qty) or qty % 1 != 0 or qty < 0:
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "need_qty должен быть "
                                    "неотрицательным целым числом."
                                )

                            if pd.isna(row["batch_id"]):
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "batch_id обязателен."
                                )

                            records.append(
                                WeldingNeed(
                                    record_id=str(row["record_id"]).strip(),
                                    date=row["date"],
                                    batch_id=str(row["batch_id"]).strip(),
                                    need_qty=int(qty),
                                )
                            )

                        WeldingNeed.objects.bulk_create(records)

                    else:
                        records = []

                        for row in df.to_dict("records"):
                            plan = pd.to_numeric(
                                row["plan_docs"], errors="coerce"
                            )
                            fact = pd.to_numeric(
                                row["fact_docs"], errors="coerce"
                            )

                            if (
                                pd.isna(plan) or pd.isna(fact)
                                or plan % 1 != 0 or fact % 1 != 0
                                or plan < 0 or fact < 0
                            ):
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "Количество документов должно быть "
                                    "неотрицательным целым числом."
                                )

                            if pd.isna(row["employee"]):
                                raise ValueError(
                                    f"{row['record_id']}: "
                                    "employee обязателен."
                                )

                            records.append(
                                DocumentationRecord(
                                    record_id=str(row["record_id"]).strip(),
                                    date=row["date"],
                                    employee=str(row["employee"]).strip(),
                                    plan_docs=int(plan),
                                    fact_docs=int(fact),
                                )
                            )

                        DocumentationRecord.objects.bulk_create(records)

                messages.success(
                    request,
                    f"Успешно импортировано строк: {len(df)}."
                )
                return redirect("upload_excel")

            except Exception as exc:
                messages.error(request, f"Импорт отменён: {exc}")

    else:
        form = ExcelUploadForm()

    return render(
        request,
        "dashboard/upload.html",
        {"form": form},
    )


def dashboard(request):
    date_from = parse_date(request.GET.get("date_from", ""))
    date_to = parse_date(request.GET.get("date_to", ""))

    stock_qs = Stock.objects.all()
    welding_qs = WeldingNeed.objects.all()
    docs_qs = DocumentationRecord.objects.all()

    if date_from:
        stock_qs = stock_qs.filter(date__gte=date_from)
        welding_qs = welding_qs.filter(date__gte=date_from)
        docs_qs = docs_qs.filter(date__gte=date_from)

    if date_to:
        stock_qs = stock_qs.filter(date__lte=date_to)
        welding_qs = welding_qs.filter(date__lte=date_to)
        docs_qs = docs_qs.filter(date__lte=date_to)

    stock = list(stock_qs)
    welding = list(welding_qs)
    docs = list(docs_qs)

    # Проверки качества данных за выбранный период.
    diagnostics = []

    # 1. Повторяющиеся идентификаторы заготовок.
    record_id_counts = Counter(item.record_id for item in stock)

    for item in stock:
        if record_id_counts[item.record_id] > 1:
            diagnostics.append({
                "severity": "error",
                "title": "Дубликат записи",
                "message": (
                    f"{item.record_id}: этот идентификатор встречается "
                    f"{record_id_counts[item.record_id]} раза."
                ),
                "data_type": "stock",
                "pk": item.pk,
            })

    # 2. Отрицательный факт.
    for item in stock:
        if item.fact_qty is not None and item.fact_qty < 0:
            diagnostics.append({
                "severity": "error",
                "title": "Отрицательное значение",
                "message": (
                    f"{item.record_id}: факт равен {item.fact_qty}."
                ),
                "data_type": "stock",
                "pk": item.pk,
            })

    # 3. Факт существенно выше плана.
    for item in stock:
        if (
            item.fact_qty is not None
            and item.plan_qty > 0
            and item.fact_qty > item.plan_qty * 2
        ):
            diagnostics.append({
                "severity": "warning",
                "title": "Проверьте фактическое значение",
                "message": (
                    f"{item.record_id}: факт {item.fact_qty} "
                    f"более чем вдвое превышает план {item.plan_qty}."
                ),
                "data_type": "stock",
                "pk": item.pk,
            })

    # 4. Отсутствующий факт.
    for item in stock:
        if item.fact_qty is None:
            diagnostics.append({
                "severity": "warning",
                "title": "Не указан факт",
                "message": (
                    f"{item.record_id}: фактическое количество не заполнено."
                ),
                "data_type": "stock",
                "pk": item.pk,
            })


    def total(items, field):
        return sum(
            getattr(item, field) or 0
            for item in items
        )

    # Сводные показатели за выбранный период.
    summary = {
        "stock_plan": total(stock, "plan_qty"),
        "stock_fact": total(stock, "fact_qty"),
        "welding_need": total(welding, "need_qty"),
        "docs_plan": total(docs, "plan_docs"),
        "docs_fact": total(docs, "fact_docs"),
    }

    # Сводка по дням.
    all_dates = sorted(
        {item.date for item in stock}
        | {item.date for item in welding}
        | {item.date for item in docs}
    )

    daily_summary = []

    for current_date in all_dates:
        day_stock = [x for x in stock if x.date == current_date]
        day_welding = [x for x in welding if x.date == current_date]
        day_docs = [x for x in docs if x.date == current_date]

        daily_summary.append({
            "date": current_date.strftime("%d.%m.%Y"),
            "stock_plan": total(day_stock, "plan_qty"),
            "stock_fact": total(day_stock, "fact_qty"),
            "welding_need": total(day_welding, "need_qty"),
            "docs_plan": total(day_docs, "plan_docs"),
            "docs_fact": total(day_docs, "fact_docs"),
        })

    # Связь заготовок и потребности сварки — по batch_id.
    batch_comparison = []

    # Считаем количество записей с каждым record_id.
    record_id_counts = Counter(item.record_id for item in stock)

    for need in welding:
        related_stock = [
            item for item in stock
            if item.batch_id == need.batch_id
        ]

        known_facts = [
            item.fact_qty for item in related_stock
            if item.fact_qty is not None
        ]

        # Проверяем качество данных, участвующих в сравнении.
        quality_issues = []

        for item in related_stock:
            if record_id_counts[item.record_id] > 1:
                quality_issues.append(
                    f"Дубликат записи {item.record_id}"
                )

            if item.fact_qty is None:
                quality_issues.append(
                    f"Не указан факт для {item.record_id}"
                )

            elif item.fact_qty < 0:
                quality_issues.append(
                    f"Отрицательный факт для {item.record_id}"
                )

            elif item.plan_qty > 0 and item.fact_qty > item.plan_qty * 2:
                quality_issues.append(
                    f"Подозрительный факт для {item.record_id}"
                )

        batch_comparison.append({
            "batch_id": need.batch_id,
            "need_qty": need.need_qty,
            "stock_records": related_stock,
            "stock_fact": sum(known_facts) if known_facts else None,
            "difference": (
                sum(known_facts) - need.need_qty
                if known_facts else None
            ),
            "quality_issues": quality_issues,
            "requires_review": bool(quality_issues),
        })

    chart_data = {
        "labels": [row["date"] for row in daily_summary],
        "stock_plan": [row["stock_plan"] for row in daily_summary],
        "stock_fact": [row["stock_fact"] for row in daily_summary],
        "welding_need": [row["welding_need"] for row in daily_summary],
        "docs_plan": [row["docs_plan"] for row in daily_summary],
        "docs_fact": [row["docs_fact"] for row in daily_summary],
    }

    context = {
        "date_from": date_from.isoformat() if date_from else "",
        "date_to": date_to.isoformat() if date_to else "",
        "summary": summary,
        "daily_summary": daily_summary,
        "batch_comparison": batch_comparison,
        "stock": stock,
        "welding": welding,
        "docs": docs,
        "chart_data": json.dumps(chart_data, ensure_ascii=False),
        "diagnostics": diagnostics,
    }

    return render(request, "dashboard/dashboard.html", context)



def _get_export_data(request):
    """Получает данные за период, указанный в фильтре панели."""
    date_from = parse_date(request.GET.get("date_from", ""))
    date_to = parse_date(request.GET.get("date_to", ""))

    stock = Stock.objects.all()
    welding = WeldingNeed.objects.all()
    docs = DocumentationRecord.objects.all()

    if date_from:
        stock = stock.filter(date__gte=date_from)
        welding = welding.filter(date__gte=date_from)
        docs = docs.filter(date__gte=date_from)

    if date_to:
        stock = stock.filter(date__lte=date_to)
        welding = welding.filter(date__lte=date_to)
        docs = docs.filter(date__lte=date_to)

    return (
        list(stock.order_by("date", "record_id")),
        list(welding.order_by("date", "record_id")),
        list(docs.order_by("date", "record_id")),
    )


def _style_sheet(sheet):
    """Оформляет заголовок и ширину столбцов."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="183153",
        )

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        sheet.column_dimensions[column[0].column_letter].width = min(
            max_length + 3, 35
        )


def export_excel(request):
    stock, welding, docs = _get_export_data(request)

    workbook = Workbook()

    # Лист 1: общие показатели.
    sheet = workbook.active
    sheet.title = "Сводка"
    sheet.append(["Показатель", "Значение"])
    sheet.append(["Заготовка — план", sum(x.plan_qty for x in stock)])
    sheet.append([
        "Заготовка — факт",
        sum(x.fact_qty or 0 for x in stock),
    ])
    sheet.append([
        "Потребность сварки",
        sum(x.need_qty for x in welding),
    ])
    sheet.append(["Документация — план", sum(x.plan_docs for x in docs)])
    sheet.append(["Документация — факт", sum(x.fact_docs for x in docs)])
    _style_sheet(sheet)

    # Лист 2: заготовка.
    sheet = workbook.create_sheet("Заготовка")
    sheet.append([
        "record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty"
    ])
    for item in stock:
        sheet.append([
            item.record_id, item.date, item.batch_id, item.employee,
            item.plan_qty, item.fact_qty,
        ])
    _style_sheet(sheet)

    # Лист 3: потребность сварки.
    sheet = workbook.create_sheet("Потребность сварки")
    sheet.append(["record_id", "date", "batch_id", "need_qty"])
    for item in welding:
        sheet.append([
            item.record_id, item.date, item.batch_id, item.need_qty
        ])
    _style_sheet(sheet)

    # Лист 4: документация.
    sheet = workbook.create_sheet("Документация")
    sheet.append([
        "record_id", "date", "employee", "plan_docs", "fact_docs"
    ])
    for item in docs:
        sheet.append([
            item.record_id, item.date, item.employee,
            item.plan_docs, item.fact_docs,
        ])
    _style_sheet(sheet)

    # Лист 5: сравнение связанных партий.
    sheet = workbook.create_sheet("Сравнение партий")
    sheet.append([
        "batch_id", "need_qty", "stock_fact", "difference"
    ])

    for need in welding:
        related = [
            item for item in stock
            if item.batch_id == need.batch_id
        ]
        known_facts = [
            item.fact_qty for item in related
            if item.fact_qty is not None
        ]

        stock_fact = sum(known_facts) if known_facts else None
        difference = (
            stock_fact - need.need_qty
            if stock_fact is not None else None
        )

        sheet.append([
            need.batch_id, need.need_qty, stock_fact, difference
        ])

    _style_sheet(sheet)

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = (
        'attachment; filename="management_report.xlsx"'
    )
    workbook.save(response)
    return response


def export_csv(request):
    stock, welding, docs = _get_export_data(request)

    response = HttpResponse(
        content_type="text/csv; charset=utf-8-sig"
    )
    response["Content-Disposition"] = (
        'attachment; filename="management_report.csv"'
    )

    writer = csv.writer(response, delimiter=";")
    writer.writerow([
        "data_type", "record_id", "date", "batch_id", "employee",
        "plan_qty", "fact_qty", "need_qty", "plan_docs", "fact_docs",
    ])

    for item in stock:
        writer.writerow([
            "stock", item.record_id, item.date, item.batch_id,
            item.employee, item.plan_qty, item.fact_qty, "", "", "",
        ])

    for item in welding:
        writer.writerow([
            "welding", item.record_id, item.date, item.batch_id,
            "", "", "", item.need_qty, "", "",
        ])

    for item in docs:
        writer.writerow([
            "documentation", item.record_id, item.date, "",
            item.employee, "", "", "", item.plan_docs, item.fact_docs,
        ])

    return response


def edit_record(request, data_type, pk):
    configurations = {
        "stock": (Stock, StockEditForm, "заготовки"),
        "welding": (WeldingNeed, WeldingNeedEditForm, "потребности сварки"),
        "documentation": (
            DocumentationRecord,
            DocumentationEditForm,
            "документации",
        ),
    }

    if data_type not in configurations:
        from django.http import Http404
        raise Http404("Неизвестный тип данных.")

    model, form_class, title = configurations[data_type]
    record = get_object_or_404(model, pk=pk)

    if request.method == "POST":
        form = form_class(request.POST, instance=record)

        if form.is_valid():
            form.save()
            messages.success(
                request,
                f"Запись {record.record_id} успешно сохранена.",
            )

            # Возвращаемся на панель с прежним фильтром дат.
            params = {}
            date_from = request.GET.get("date_from", "")
            date_to = request.GET.get("date_to", "")

            if date_from:
                params["date_from"] = date_from
            if date_to:
                params["date_to"] = date_to

            query_string = urlencode(params)
            return redirect("/" + (f"?{query_string}" if query_string else ""))
    else:
        form = form_class(instance=record)

    return render(
        request,
        "dashboard/edit_record.html",
        {
            "form": form,
            "record": record,
            "title": title,
            "data_type": data_type,
            "date_from": request.GET.get("date_from", ""),
            "date_to": request.GET.get("date_to", ""),
        },
    )
